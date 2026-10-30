"""Создание Word-протоколов: по шаблону с метками {{...}} или встроенному образцу."""
import copy
import datetime
import io
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt

from .calcd import fmt

POINT_MARK = "{{т."  # строка таблицы с такими метками повторяется для каждой точки


def channel_fields(ch, digits=4):
    f = lambda v: fmt(v, digits)
    r = ch.result
    return {
        "№": ch.num, "KKS": ch.kks, "НАИМЕНОВАНИЕ": ch.name,
        "НИЖН": fmt(ch.low), "ВЕРХН": fmt(ch.high), "ЕД": ch.unit,
        "ДИАПАЗОН": f"{fmt(ch.low)} … {fmt(ch.high)} {ch.unit}".strip(),
        "ПИП": ch.pip, "ПОГР_ПИП": f(ch.pip_value), "МОДУЛЬ": ch.module,
        "ПОГР_МОДУЛЯ": f(ch.module_err), "ДОПУСТ": f(ch.allowed_value), "ЗАМЕТКИ": ch.notes,
        "a1": f(r.a1) if r else "", "a2": f(r.a2) if r else "",
        "b1": f(r.b1) if r else "", "b2": f(r.b2) if r else "",
        "gh": f(r.gh) if r else "", "gl": f(r.gl) if r else "",
        "ЭТ_ИК": f(ch.et_ik), "ОБЩАЯ": f(ch.total), "РЕЗУЛЬТАТ": ch.verdict,
        "ДАТА": datetime.date.today().strftime("%d.%m.%Y"),
    }


def point_fields(ch, i, digits=4):
    f = lambda v: fmt(v, digits)
    p = ch.points[i]
    rp = ch.result.points[i] if ch.result else None
    am = sum(p.measured) / len(p.measured) if p.measured else None
    return {
        "т.№": str(i + 1), "т.%": fmt(None if p.pct is None else round(p.pct, 2)), "т.вход": fmt(p.input),
        "т.xd": fmt(p.xd), "т.изм": fmt(am),
        "т.ds": f(rp.ds) if rp else "", "т.dh": f(rp.dh) if rp else "",
        "т.dl": f(rp.dl) if rp else "",
        "т.верх": f(rp.upper) if rp else "", "т.низ": f(rp.lower) if rp else "",
    }


def _replace_in_paragraph(par, values):
    text = "".join(r.text for r in par.runs)
    if "{{" not in text:
        return
    new = re.sub(r"\{\{\s*([^{}]+?)\s*\}\}", lambda m: values.get(m.group(1), m.group(0)), text)
    if new == text:
        return
    # метка может быть разбита Word'ом на несколько фрагментов: текст в первый
    par.runs[0].text = new
    for r in par.runs[1:]:
        r.text = ""


def _iter_paragraphs(container):
    for p in container.paragraphs:
        yield p
    for t in container.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MARK_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def _fill_xml(el, values):
    """Подставить метки во все абзацы элемента (работа прямо с XML - быстро)."""
    for p in el.iter(W + "p"):
        ts = list(p.iter(W + "t"))
        text = "".join(t.text or "" for t in ts)
        if "{{" not in text:
            continue
        new = MARK_RE.sub(lambda m: values.get(m.group(1), m.group(0)), text)
        if new == text:
            continue
        # метка может быть разбита Word'ом на несколько фрагментов: текст в первый
        ts[0].text = new
        ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        for t in ts[1:]:
            t.text = ""


def _xml_text(el):
    return "".join(t.text or "" for t in el.iter(W + "t"))


def _fill_protocol(elements, ch, digits):
    """Строки таблицы с метками точки размножить по точкам, затем общие метки."""
    for el in elements:
        for tr in [tr for tr in el.iter(W + "tr") if POINT_MARK in _xml_text(tr)]:
            for i in range(len(ch.points)):
                new_tr = copy.deepcopy(tr)
                _fill_xml(new_tr, point_fields(ch, i, digits))
                tr.addprevious(new_tr)
            tr.getparent().remove(tr)
        _fill_xml(el, channel_fields(ch, digits))


def _page_break_before(el):
    if el.tag != W + "p":
        return
    ppr = el.find(W + "pPr")
    if ppr is None:
        ppr = el.makeelement(W + "pPr", {})
        el.insert(0, ppr)
    if ppr.find(W + "pageBreakBefore") is None:
        ppr.insert(0, ppr.makeelement(W + "pageBreakBefore", {}))


def from_template(template, ch, digits=4):
    return build([ch], template, digits)


def _cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=None):
    cell.text = ""
    par = cell.paragraphs[0]
    par.alignment = align
    run = par.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)


def _table(doc, rows, cols, widths):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Cm(w)
    return t


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text).bold = True


def make_template():
    """Образец протокола на одну страницу A4 с метками {{...}}."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.0)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    C = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = C
    r = p.add_run("ПРОТОКОЛ № {{№}}")
    r.bold = True
    r.font.size = Pt(14)
    p = doc.add_paragraph()
    p.alignment = C
    p.add_run("аттестации измерительного канала").bold = True
    p = doc.add_paragraph("Дата: {{ДАТА}}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(6)

    info = (("Позиционное обозначение (KKS)", "{{KKS}}"),
            ("Наименование ИК", "{{НАИМЕНОВАНИЕ}}"),
            ("Диапазон измерения ИК", "{{НИЖН}} … {{ВЕРХН}} {{ЕД}}"),
            ("Первичный измерительный преобразователь (ПИП)", "{{ПИП}}"),
            ("Модуль ввода", "{{МОДУЛЬ}}"),
            ("Погрешность модуля, %", "{{ПОГР_МОДУЛЯ}}"))
    t = _table(doc, len(info), 2, (7.0, 10.5))
    for row, (k, v) in zip(t.rows, info):
        _cell(row.cells[0], k)
        _cell(row.cells[1], v, bold=k.startswith("Позиц"))

    _heading(doc, "1. Результаты измерений в контрольных точках")
    head = ("№", "КТ, %", "Действит. значение, {{ЕД}}", "Измеренное значение, {{ЕД}}",
            "Систем. составл. ds", "Граница Δ верх.", "Граница Δ нижн.")
    marks = ("{{т.№}}", "{{т.%}}", "{{т.xd}}", "{{т.изм}}", "{{т.ds}}", "{{т.dh}}", "{{т.dl}}")
    t = _table(doc, 2, len(head), (0.9, 1.4, 3.0, 3.0, 3.0, 3.1, 3.1))
    for c, h in zip(t.rows[0].cells, head):
        _cell(c, h, bold=True, align=C, size=10)
    for c, m in zip(t.rows[1].cells, marks):
        _cell(c, m, align=C)

    _heading(doc, "2. Коэффициенты (расчёт по методике CalcD)")
    coef = (("a1", "{{a1}}", "a2", "{{a2}}"), ("b1", "{{b1}}", "b2", "{{b2}}"),
            ("gh, %", "{{gh}}", "gl, %", "{{gl}}"))
    t = _table(doc, 3, 4, (2.0, 6.75, 2.0, 6.75))
    for row, vals in zip(t.rows, coef):
        for i, (c, v) in enumerate(zip(row.cells, vals)):
            _cell(c, v, bold=i % 2 == 0, align=C)

    _heading(doc, "3. Приведённая погрешность ИК")
    err = (("Погрешность ПИП, %", "{{ПОГР_ПИП}}"),
           ("Погрешность ЭТ ИК, % (max |gh|, |gl|)", "{{ЭТ_ИК}}"),
           ("Общая погрешность ИК, % (√(ПИП² + ЭТ ИК²))", "{{ОБЩАЯ}}"),
           ("Допустимая погрешность ИК, %", "{{ДОПУСТ}}"))
    t = _table(doc, len(err), 2, (11.5, 6.0))
    for row, (k, v) in zip(t.rows, err):
        _cell(row.cells[0], k)
        _cell(row.cells[1], v, align=C)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.add_run("Заключение: ").bold = True
    p.add_run("{{РЕЗУЛЬТАТ}}").bold = True
    p = doc.add_paragraph("Примечание: {{ЗАМЕТКИ}}")
    p.paragraph_format.space_before = Pt(4)

    for who in ("Аттестацию провёл", "Проверил"):
        p = doc.add_paragraph(f"{who}:  ____________________  /____________________/      «___» __________ 20__ г.")
        p.paragraph_format.space_before = Pt(18)
    return doc


def default_template():
    buf = io.BytesIO()
    make_template().save(buf)
    buf.seek(0)
    return buf


def build(channels, template=None, digits=4):
    """Один документ: протоколы каналов подряд, каждый с новой страницы.

    Шаблон открывается один раз, тело копируется на каждый канал - так
    3000 протоколов собираются за секунды, а не минуты.
    """
    doc = Document(template or default_template())
    body = doc.element.body
    sect = body.sectPr
    proto = [el for el in body.iterchildren() if el is not sect]
    for el in proto:
        body.remove(el)
    for n, ch in enumerate(channels):
        els = [copy.deepcopy(el) for el in proto]
        _fill_protocol(els, ch, digits)
        if n and els:
            _page_break_before(els[0])
        for el in els:
            if sect is not None:
                sect.addprevious(el)
            else:
                body.append(el)
    # колонтитулы общие для документа - метки из первого канала
    if channels:
        vals = channel_fields(channels[0], digits)
        for sec in doc.sections:
            for part in (sec.header, sec.footer):
                for p in _iter_paragraphs(part):
                    _replace_in_paragraph(p, vals)
    return doc


def safe_name(s):
    s = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", s).strip(" .")
    return s or "канал"


def generate(channels, out_dir, template=None, digits=4):
    """Создать один .docx со всеми протоколами. Возвращает путь к файлу."""
    os.makedirs(out_dir, exist_ok=True)
    first, last = channels[0], channels[-1]
    name = "Протоколы"
    if first.num and last.num:
        name += f" №{first.num}" + (f"-{last.num}" if last is not first else "")
    elif first.kks:
        name += f" {first.kks}" + (f" - {last.kks}" if last is not first else "")
    base = os.path.join(out_dir, safe_name(name))
    path, n = base + ".docx", 2
    while os.path.exists(path):
        path, n = f"{base} ({n}).docx", n + 1
    build(channels, template, digits).save(path)
    return path
